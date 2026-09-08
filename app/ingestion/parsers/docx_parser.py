from pathlib import Path

from docx import Document

from .base import BaseDocumentParser
from ..models.document import ParsedDocument, DocumentSection


class DocxParser(BaseDocumentParser):

    def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"DOCX file not found: {file_path}"
            )

        if path.suffix.lower() != ".docx":
            raise ValueError(
                f"Expected a DOCX file, got: {path.suffix}"
            )

        document = Document(file_path)

        sections = []
        current_title = "General"
        current_content = []
        section_counter = 1

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if not text:
                continue

            style_name = paragraph.style.name.lower()

            if style_name.startswith("heading"):
                if current_content:
                    sections.append(
                        DocumentSection(
                            section_id=f"SEC-{section_counter:03d}",
                            title=current_title,
                            content="\n".join(current_content),
                        )
                    )

                    section_counter += 1
                    current_content = []

                current_title = text

            else:
                current_content.append(text)

        if current_content:
            sections.append(
                DocumentSection(
                    section_id=f"SEC-{section_counter:03d}",
                    title=current_title,
                    content="\n".join(current_content),
                )
            )

        full_text_parts = []

        for section in sections:
            full_text_parts.append(
                f"{section.title}\n{section.content}"
            )

        full_text = "\n\n".join(full_text_parts)

        return ParsedDocument(
            document_name=path.name,
            document_type="docx",
            source_path=str(path),
            full_text=full_text,
            sections=sections,
        )
