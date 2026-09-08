from pathlib import Path
from openpyxl.styles import PatternFill
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from pydantic import BaseModel

from app.analysis.models.requirement import RequirementAnalysis
from app.strategy.models.test_strategy import TestStrategy
from app.scenario.models.test_scenario import ScenarioAnalysis
from app.test_design.models.test_design import TestDesignAnalysis
from app.test_case.models.test_case import TestCaseAnalysis
from app.rtm.models.rtm import RTMAnalysis


class ExcelGeneratorInput(BaseModel):
    requirement_analysis: RequirementAnalysis
    test_strategy: TestStrategy
    scenario_analysis: ScenarioAnalysis
    test_design_analysis: TestDesignAnalysis
    test_case_analysis: TestCaseAnalysis
    rtm_analysis: RTMAnalysis


class ExcelGenerator:
    """Deterministic Excel workbook generator for validated QA artifacts."""

    SHEET_NAMES = [
        "Summary",
        "Requirements",
        "Risk Strategy",
        "Scenarios",
        "Test Designs",
        "Test Cases",
        "RTM",
    ]

    def generate(
        self,
        data: ExcelGeneratorInput,
        output_path: str | Path,
    ) -> Path:
        output_path = Path(output_path)

        workbook = Workbook()
        workbook.remove(workbook.active)

        for sheet_name in self.SHEET_NAMES:
            workbook.create_sheet(sheet_name)

        self._populate_summary(workbook, data)
        self._populate_requirements(workbook, data)
        self._populate_risk_strategy(workbook, data)
        self._populate_scenarios(workbook, data)
        self._populate_test_designs(workbook, data)
        self._populate_test_cases(workbook, data)
        self._populate_rtm(workbook, data)
        for sheet in workbook.worksheets:
            self._format_sheet(sheet)
        self._apply_status_formatting(workbook)
        self._add_excel_tables(workbook)
        self._apply_number_formatting(workbook)
        self._format_summary(workbook)

        workbook.save(output_path)

        return output_path

    def _populate_summary(
        self,
        workbook: Workbook,
        data: ExcelGeneratorInput,
    ) -> None:
        sheet = workbook["Summary"]

        sheet.append(["Metric", "Value"])

        sheet.append([
            "Document Name",
            data.requirement_analysis.document_name,
        ])

        sheet.append([
            "Total Requirements",
            len(data.requirement_analysis.requirements),
        ])

        sheet.append([
            "Total Risks",
            len(data.test_strategy.requirement_risks),
        ])

        sheet.append([
            "Total Scenarios",
            len(data.scenario_analysis.scenarios),
        ])

        sheet.append([
            "Total Test Designs",
            len(data.test_design_analysis.test_designs),
        ])

        sheet.append([
            "Total Test Cases",
            len(data.test_case_analysis.test_cases),
        ])

        sheet.append([
            "Covered Requirements",
            data.rtm_analysis.covered_requirements,
        ])

        sheet.append([
            "Partially Covered Requirements",
            data.rtm_analysis.partially_covered_requirements,
        ])

        sheet.append([
            "Not Covered Requirements",
            data.rtm_analysis.not_covered_requirements,
        ])

        sheet.append([
            "Overall Coverage %",
            data.rtm_analysis.overall_coverage_percentage / 100,
        ])

        sheet.append([
            "Overall Risk Level",
            data.test_strategy.overall_risk_level.value,
        ])

        sheet.append([
            "Requirement Confidence",
            data.requirement_analysis.overall_confidence,
        ])

        sheet.append([
            "Strategy Confidence",
            data.test_strategy.overall_strategy_confidence,
        ])

        sheet.append([
            "Scenario Confidence",
            data.scenario_analysis.overall_scenario_confidence,
        ])

        sheet.append([
            "Test Design Confidence",
            data.test_design_analysis.overall_design_confidence,
        ])

        sheet.append([
            "Test Case Confidence",
            data.test_case_analysis.overall_test_case_confidence,
        ])

    def _populate_requirements(
            self,
            workbook: Workbook,
            data: ExcelGeneratorInput,
        ) -> None:
            sheet = workbook["Requirements"]

            headers = [
                "Requirement ID",
                "Title",
                "Description",
                "Requirement Type",
                "Classification",
                "Source Document",
                "Source Section",
                "Source Page",
                "Source Text",
                "Business Rules",
                "Dependencies",
                "Confidence",
            ]

            sheet.append(headers)

            for requirement in data.requirement_analysis.requirements:
                sheet.append([
                    requirement.requirement_id,
                    requirement.title,
                    requirement.description,
                    requirement.requirement_type.value,
                    requirement.classification.value,
                    requirement.source_document,
                    requirement.source_section,
                    requirement.source_page,
                    requirement.source_text,
                    "\n".join(requirement.business_rules),
                    "\n".join(requirement.dependencies),
                    requirement.confidence,
                ])

    def _populate_risk_strategy(
        self,
        workbook: Workbook,
        data: ExcelGeneratorInput,
    ) -> None:
        sheet = workbook["Risk Strategy"]

        headers = [
            "Requirement ID",
            "Risk Level",
            "Risk Score",
            "Risk Reason",
            "Impacted Areas",
            "Recommended Test Types",
        ]

        sheet.append(headers)

        for risk in data.test_strategy.requirement_risks:
            sheet.append([
                risk.requirement_id,
                risk.risk_level.value,
                risk.risk_score,
                risk.risk_reason,
                "\n".join(risk.impacted_areas),
                "\n".join(
                    test_type.value
                    for test_type in risk.recommended_test_types
                ),
            ])

    def _populate_scenarios(
        self,
        workbook: Workbook,
        data: ExcelGeneratorInput,
    ) -> None:
        sheet = workbook["Scenarios"]

        headers = [
            "Scenario ID",
            "Requirement ID",
            "Title",
            "Description",
            "Scenario Type",
            "Priority",
            "Preconditions",
            "Expected Behavior",
            "Source Requirement",
        ]

        sheet.append(headers)

        for scenario in data.scenario_analysis.scenarios:
            sheet.append([
                scenario.scenario_id,
                scenario.requirement_id,
                scenario.title,
                scenario.description,
                scenario.scenario_type.value,
                scenario.priority,
                "\n".join(scenario.preconditions),
                scenario.expected_behavior,
                scenario.source_requirement,
            ])

    def _populate_test_designs(
        self,
        workbook: Workbook,
        data: ExcelGeneratorInput,
    ) -> None:
        sheet = workbook["Test Designs"]

        headers = [
            "Design ID",
            "Scenario ID",
            "Requirement ID",
            "Title",
            "Objective",
            "Test Design Technique",
            "Test Data Conditions",
            "Coverage Area",
            "Expected Focus",
            "Source Scenario",
        ]

        sheet.append(headers)

        for design in data.test_design_analysis.test_designs:
            sheet.append([
                design.design_id,
                design.scenario_id,
                design.requirement_id,
                design.title,
                design.objective,
                design.test_design_technique.value,
                "\n".join(design.test_data_conditions),
                design.coverage_area,
                design.expected_focus,
                design.source_scenario,
            ])
    def _populate_test_cases(
        self,
        workbook: Workbook,
        data: ExcelGeneratorInput,
    ) -> None:
        sheet = workbook["Test Cases"]

        headers = [
            "Test Case ID",
            "Design ID",
            "Scenario ID",
            "Requirement ID",
            "Title",
            "Objective",
            "Test Type",
            "Priority",
            "Preconditions",
            "Test Data",
            "Steps",
            "Expected Results",
            "Source Design",
        ]

        sheet.append(headers)

        for test_case in data.test_case_analysis.test_cases:
            sheet.append([
                test_case.test_case_id,
                test_case.design_id,
                test_case.scenario_id,
                test_case.requirement_id,
                test_case.title,
                test_case.objective,
                test_case.test_type.value,
                test_case.priority,
                "\n".join(test_case.preconditions),
                "\n".join(test_case.test_data),
                "\n".join(test_case.steps),
                "\n".join(test_case.expected_results),
                test_case.source_design,
            ])
    def _populate_rtm(
        self,
        workbook: Workbook,
        data: ExcelGeneratorInput,
    ) -> None:
        sheet = workbook["RTM"]

        headers = [
            "Requirement ID",
            "Scenario IDs",
            "Design IDs",
            "Test Case IDs",
            "Coverage Status",
            "Coverage Notes",
        ]

        sheet.append(headers)

        for coverage in data.rtm_analysis.requirement_coverage:
            sheet.append([
                coverage.requirement_id,
                "\n".join(coverage.scenario_ids),
                "\n".join(coverage.design_ids),
                "\n".join(coverage.test_case_ids),
                coverage.coverage_status.value,
                coverage.coverage_notes,
            ])
    def _format_sheet(self, sheet) -> None:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = cell.font.copy(bold=True)

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = cell.alignment.copy(
                    wrap_text=True,
                    vertical="top",
                )

        for column_cells in sheet.columns:
            max_length = 0

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )

            column_letter = column_cells[0].column_letter
            sheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                50,
            )
    def _apply_status_formatting(self, workbook: Workbook) -> None:
        risk_sheet = workbook["Risk Strategy"]

        risk_fills = {
            "Low": PatternFill(fill_type="solid", fgColor="C6EFCE"),
            "Medium": PatternFill(fill_type="solid", fgColor="FFEB9C"),
            "High": PatternFill(fill_type="solid", fgColor="F4B084"),
            "Critical": PatternFill(fill_type="solid", fgColor="F4CCCC"),
        }

        for row in risk_sheet.iter_rows(min_row=2):
            risk_cell = row[1]
            fill = risk_fills.get(risk_cell.value)

            if fill:
                risk_cell.fill = fill

        rtm_sheet = workbook["RTM"]

        coverage_fills = {
            "Covered": PatternFill(fill_type="solid", fgColor="C6EFCE"),
            "Partial": PatternFill(fill_type="solid", fgColor="FFEB9C"),
            "Not Covered": PatternFill(fill_type="solid", fgColor="F4CCCC"),
        }

        for row in rtm_sheet.iter_rows(min_row=2):
            coverage_cell = row[4]
            fill = coverage_fills.get(coverage_cell.value)

            if fill:
                coverage_cell.fill = fill
    def _add_excel_tables(self, workbook: Workbook) -> None:
        table_sheets = [
            "Requirements",
            "Risk Strategy",
            "Scenarios",
            "Test Designs",
            "Test Cases",
            "RTM",
        ]

        for index, sheet_name in enumerate(table_sheets, start=1):
            sheet = workbook[sheet_name]

            if sheet.max_row < 2:
                continue

            table_name = f"QA_Table_{index}"

            table = Table(
                displayName=table_name,
                ref=sheet.dimensions,
            )

            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = style
            sheet.add_table(table)
    def _apply_number_formatting(self, workbook: Workbook) -> None:
        summary = workbook["Summary"]

        for row in summary.iter_rows(min_row=2):
            metric = row[0].value

            if metric == "Overall Coverage %":
                row[1].number_format = "0.0%"

            elif metric.endswith("Confidence"):
                row[1].number_format = "0.0%"
    def _format_summary(self, workbook: Workbook) -> None:
        sheet = workbook["Summary"]

        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 24

        for cell in sheet[1]:
            cell.font = cell.font.copy(bold=True, size=12)

        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=1).font = sheet.cell(
                row=row,
                column=1,
            ).font.copy(bold=True)
