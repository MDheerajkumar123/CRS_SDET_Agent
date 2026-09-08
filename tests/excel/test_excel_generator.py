from pathlib import Path

from openpyxl import load_workbook

from app.analysis.models.requirement import (
    RequirementAnalysis,
    Requirement,
    RequirementClassification,
    RequirementType,
)
from app.strategy.models.test_strategy import (
    TestStrategy,
    RequirementRisk,
    RiskLevel,
)
from app.scenario.models.test_scenario import (
    ScenarioAnalysis,
    TestScenario,
    ScenarioType,
)
from app.test_design.models.test_design import (
    TestDesignAnalysis,
    TestDesign,
    TestDesignTechnique,
)
from app.test_case.models.test_case import (
    TestCaseAnalysis,
    TestCase,
    TestCaseType,
)
from app.rtm.models.rtm import (
    RTMAnalysis,
    RequirementCoverage,
    CoverageStatus,
)

from app.excel.excel_generator import ExcelGenerator, ExcelGeneratorInput


def create_test_input() -> ExcelGeneratorInput:
    requirement = Requirement(
        requirement_id="REQ-001",
        title="User Login",
        description="User shall be able to log in.",
        requirement_type=RequirementType.FUNCTIONAL,
        classification=RequirementClassification.EXPLICIT,
        source_document="sample_crs.docx",
        source_section="Authentication",
        source_page=1,
        source_text="User shall be able to log in.",
        confidence=0.95,
    )

    risk = RequirementRisk(
        requirement_id="REQ-001",
        risk_level=RiskLevel.HIGH,
        risk_score=80,
        risk_reason="Authentication is business critical.",
        impacted_areas=["Login"],
    )

    scenario = TestScenario(
        scenario_id="SCN-001",
        requirement_id="REQ-001",
        title="Successful Login",
        description="Verify valid credentials allow login.",
        scenario_type=ScenarioType.POSITIVE,
        priority="High",
        expected_behavior="User is successfully authenticated.",
        source_requirement="REQ-001",
    )

    design = TestDesign(
        design_id="TD-001",
        scenario_id="SCN-001",
        requirement_id="REQ-001",
        title="Valid Credential Design",
        objective="Validate successful login.",
        test_design_technique=TestDesignTechnique.USE_CASE,
        coverage_area="Authentication",
        expected_focus="Successful authentication",
        source_scenario="SCN-001",
    )

    test_case = TestCase(
        test_case_id="TC-001",
        design_id="TD-001",
        scenario_id="SCN-001",
        requirement_id="REQ-001",
        title="Login with Valid Credentials",
        objective="Verify successful login.",
        test_type=TestCaseType.FUNCTIONAL,
        priority="High",
        steps=["Enter valid username", "Enter valid password", "Click Login"],
        expected_results=["User is logged in successfully."],
        source_design="TD-001",
    )

    coverage = RequirementCoverage(
        requirement_id="REQ-001",
        scenario_ids=["SCN-001"],
        design_ids=["TD-001"],
        test_case_ids=["TC-001"],
        coverage_status=CoverageStatus.COVERED,
        coverage_notes="Fully covered.",
    )

    return ExcelGeneratorInput(
        requirement_analysis=RequirementAnalysis(
            document_name="sample_crs.docx",
            analysis_summary="Sample analysis.",
            requirements=[requirement],
            overall_confidence=0.95,
        ),
        test_strategy=TestStrategy(
            document_name="sample_crs.docx",
            strategy_summary="Sample strategy.",
            requirement_risks=[risk],
            prioritized_requirements=["REQ-001"],
            overall_risk_level=RiskLevel.HIGH,
            overall_strategy_confidence=0.90,
        ),
        scenario_analysis=ScenarioAnalysis(
            document_name="sample_crs.docx",
            scenario_summary="Sample scenarios.",
            scenarios=[scenario],
            overall_scenario_confidence=0.92,
        ),
        test_design_analysis=TestDesignAnalysis(
            document_name="sample_crs.docx",
            design_summary="Sample designs.",
            test_designs=[design],
            overall_design_confidence=0.91,
        ),
        test_case_analysis=TestCaseAnalysis(
            document_name="sample_crs.docx",
            test_case_summary="Sample test cases.",
            test_cases=[test_case],
            overall_test_case_confidence=0.93,
        ),
        rtm_analysis=RTMAnalysis(
            document_name="sample_crs.docx",
            rtm_summary="Sample RTM.",
            requirement_coverage=[coverage],
            total_requirements=1,
            covered_requirements=1,
            partially_covered_requirements=0,
            not_covered_requirements=0,
            overall_coverage_percentage=100.0,
        ),
    )


def test_generate_creates_required_workbook_sheets(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_output.xlsx"

    result = ExcelGenerator().generate(data, output_path)

    assert result == Path(output_path)
    assert output_path.exists()

    workbook = load_workbook(output_path)

    assert workbook.sheetnames == [
        "Summary",
        "Requirements",
        "Risk Strategy",
        "Scenarios",
        "Test Designs",
        "Test Cases",
        "RTM",
    ]


def test_summary_contains_expected_metrics(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_summary.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Summary"]

    summary = {
        row[0].value: row[1].value
        for row in sheet.iter_rows(min_row=2)
    }

    assert summary["Document Name"] == "sample_crs.docx"
    assert summary["Total Requirements"] == 1
    assert summary["Total Risks"] == 1
    assert summary["Total Scenarios"] == 1
    assert summary["Total Test Designs"] == 1
    assert summary["Total Test Cases"] == 1
    assert summary["Covered Requirements"] == 1
    assert summary["Partially Covered Requirements"] == 0
    assert summary["Not Covered Requirements"] == 0
    assert summary["Overall Coverage %"] == 1.0
    assert summary["Overall Risk Level"] == "High"

def test_requirements_sheet_contains_expected_data(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_requirements.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Requirements"]

    headers = [cell.value for cell in sheet[1]]

    assert headers == [
        "Requirement ID",
        "Title",
        "Description",
        "Requirement Type",
        "Classification",
        "Source Document",
        "Source Section",
        "Source Page",
        "Source Text",
        "Business Rules",
        "Dependencies",
        "Confidence",
    ]

    assert sheet.max_row == 2

    assert sheet["A2"].value == "REQ-001"
    assert sheet["B2"].value == "User Login"
    assert sheet["D2"].value == "Functional"
    assert sheet["E2"].value == "Explicit"
    assert sheet["F2"].value == "sample_crs.docx"
    assert sheet["G2"].value == "Authentication"
    assert sheet["H2"].value == 1
    assert sheet["L2"].value == 0.95

def test_risk_strategy_sheet_contains_expected_data(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_risk_strategy.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Risk Strategy"]

    headers = [cell.value for cell in sheet[1]]

    assert headers == [
        "Requirement ID",
        "Risk Level",
        "Risk Score",
        "Risk Reason",
        "Impacted Areas",
        "Recommended Test Types",
    ]

    assert sheet.max_row == 2

    assert sheet["A2"].value == "REQ-001"
    assert sheet["B2"].value == "High"
    assert sheet["C2"].value == 80
    assert sheet["D2"].value == "Authentication is business critical."
    assert sheet["E2"].value == "Login"

def test_scenarios_sheet_contains_expected_data(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_scenarios.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Scenarios"]

    headers = [cell.value for cell in sheet[1]]

    assert headers == [
        "Scenario ID",
        "Requirement ID",
        "Title",
        "Description",
        "Scenario Type",
        "Priority",
        "Preconditions",
        "Expected Behavior",
        "Source Requirement",
    ]

    assert sheet.max_row == 2

    assert sheet["A2"].value == "SCN-001"
    assert sheet["B2"].value == "REQ-001"
    assert sheet["C2"].value == "Successful Login"
    assert sheet["D2"].value == "Verify valid credentials allow login."
    assert sheet["E2"].value == "Positive"
    assert sheet["F2"].value == "High"
    assert sheet["H2"].value == "User is successfully authenticated."
    assert sheet["I2"].value == "REQ-001"

def test_test_designs_sheet_contains_expected_data(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_test_designs.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Test Designs"]

    headers = [cell.value for cell in sheet[1]]

    assert headers == [
        "Design ID",
        "Scenario ID",
        "Requirement ID",
        "Title",
        "Objective",
        "Test Design Technique",
        "Test Data Conditions",
        "Coverage Area",
        "Expected Focus",
        "Source Scenario",
    ]

    assert sheet.max_row == 2

    assert sheet["A2"].value == "TD-001"
    assert sheet["B2"].value == "SCN-001"
    assert sheet["C2"].value == "REQ-001"
    assert sheet["D2"].value == "Valid Credential Design"
    assert sheet["E2"].value == "Validate successful login."
    assert sheet["F2"].value == "Use Case"
    assert sheet["H2"].value == "Authentication"
    assert sheet["I2"].value == "Successful authentication"
    assert sheet["J2"].value == "SCN-001"
def test_test_cases_sheet_contains_expected_data(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_test_cases.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Test Cases"]

    headers = [cell.value for cell in sheet[1]]

    assert headers == [
        "Test Case ID",
        "Design ID",
        "Scenario ID",
        "Requirement ID",
        "Title",
        "Objective",
        "Test Type",
        "Priority",
        "Preconditions",
        "Test Data",
        "Steps",
        "Expected Results",
        "Source Design",
    ]

    assert sheet.max_row == 2

    assert sheet["A2"].value == "TC-001"
    assert sheet["B2"].value == "TD-001"
    assert sheet["C2"].value == "SCN-001"
    assert sheet["D2"].value == "REQ-001"
    assert sheet["E2"].value == "Login with Valid Credentials"
    assert sheet["F2"].value == "Verify successful login."
    assert sheet["G2"].value == "Functional"
    assert sheet["H2"].value == "High"
    assert sheet["K2"].value == (
        "Enter valid username\n"
        "Enter valid password\n"
        "Click Login"
    )
    assert sheet["L2"].value == "User is logged in successfully."
    assert sheet["M2"].value == "TD-001"
def test_rtm_sheet_contains_expected_traceability(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_rtm.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["RTM"]

    headers = [cell.value for cell in sheet[1]]

    assert headers == [
        "Requirement ID",
        "Scenario IDs",
        "Design IDs",
        "Test Case IDs",
        "Coverage Status",
        "Coverage Notes",
    ]

    assert sheet.max_row == 2

    assert sheet["A2"].value == "REQ-001"
    assert sheet["B2"].value == "SCN-001"
    assert sheet["C2"].value == "TD-001"
    assert sheet["D2"].value == "TC-001"
    assert sheet["E2"].value == "Covered"
    assert sheet["F2"].value == "Fully covered."
def test_workbook_sheets_have_basic_formatting(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_formatted.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)

    for sheet in workbook.worksheets:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == sheet.dimensions

        for cell in sheet[1]:
            assert cell.font.bold is True

    requirements = workbook["Requirements"]

    assert requirements["C2"].alignment.wrap_text is True
    assert requirements.column_dimensions["C"].width <= 50
def test_risk_and_coverage_status_formatting(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_status_formatting.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)

    risk_sheet = workbook["Risk Strategy"]
    rtm_sheet = workbook["RTM"]

    assert risk_sheet["B2"].value == "High"
    assert risk_sheet["B2"].fill.fill_type == "solid"
    assert risk_sheet["B2"].fill.fgColor.rgb == "00F4B084"

    assert rtm_sheet["E2"].value == "Covered"
    assert rtm_sheet["E2"].fill.fill_type == "solid"
    assert rtm_sheet["E2"].fill.fgColor.rgb == "00C6EFCE"
def test_artifact_sheets_have_excel_tables(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_tables.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)

    artifact_sheets = [
        "Requirements",
        "Risk Strategy",
        "Scenarios",
        "Test Designs",
        "Test Cases",
        "RTM",
    ]

    for sheet_name in artifact_sheets:
        sheet = workbook[sheet_name]

        assert len(sheet.tables) == 1

        table = next(iter(sheet.tables.values()))

        assert table.ref == sheet.dimensions
def test_summary_percentage_number_formats(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_percentage_format.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Summary"]

    metrics = {
        row[0].value: row[1]
        for row in sheet.iter_rows(min_row=2)
    }

    assert metrics["Overall Coverage %"].value == 1.0
    assert metrics["Overall Coverage %"].number_format == "0.0%"

    assert metrics["Requirement Confidence"].value == 0.95
    assert metrics["Requirement Confidence"].number_format == "0.0%"

    assert metrics["Strategy Confidence"].value == 0.90
    assert metrics["Strategy Confidence"].number_format == "0.0%"
def test_summary_has_expected_styling(tmp_path):
    data = create_test_input()

    output_path = tmp_path / "qa_summary_styling.xlsx"

    ExcelGenerator().generate(data, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Summary"]

    assert sheet.column_dimensions["A"].width == 32
    assert sheet.column_dimensions["B"].width == 24

    assert sheet["A1"].font.bold is True
    assert sheet["A1"].font.sz == 12

    assert sheet["A2"].font.bold is True
    assert sheet["A3"].font.bold is True
